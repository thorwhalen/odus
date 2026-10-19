"""Plotting helpers for trajectories and matrices.

Heatmaps, life-course plots, iteration over a store's trajectory plots, and
writing a selection of trajectories out to a (multi-page) file.

``plot_trajectory`` is the readable rendering of a single trajectory: a gap
between cells (wider every ``major_every`` years), alternating grey rows so the
eye can follow a row back to its variable, optional light shading of variable
groups, and per-variable colors (for instance the ones read from the source
spreadsheet with ``row_colors_of_xls``).
"""

import matplotlib.pylab as plt
from matplotlib.colors import to_rgb, is_color_like
from itertools import islice, cycle
from collections.abc import Iterable, Mapping
import numpy as np
import pandas as pd

from odus.util import write_images

DFLT_FG_COLOR = "#1b2a49"  # color of a cell whose value is "on" (max value)
DFLT_ROW_BACKGROUNDS = ("#ffffff", "#ebebeb")  # alternated, row by row
DFLT_GROUP_COLORS = ("#9ecae1", "#fdae6b", "#a1d99b", "#bcbddc")  # light, distinct
DFLT_GROUP_TINT = 0.35  # how much a group color tints the row background
DFLT_CELL_SIZE = 0.22  # inches per cell, when the figure size is computed
DFLT_MAJOR_EVERY = 5  # a wider gap before every age/year multiple of this

ihead = lambda it: islice(it, 0, 5)  # just a little useful util. Not used within module


def heatmap(
    X,
    y=None,
    col_labels=None,
    figsize=None,
    cmap=None,
    return_gcf=False,
    ax=None,
    xlabel_top=True,
    ylabel_left=True,
    xlabel_bottom=True,
    ylabel_right=True,
    **kwargs,
):
    n_items, n_cols = X.shape
    if col_labels is not None:
        if col_labels is not False:
            assert len(col_labels) == n_cols, (
                "col_labels length should be the same as the number of columns in the matrix"
            )
    elif isinstance(X, pd.DataFrame):
        col_labels = list(X.columns)

    if figsize is None:
        x_size, y_size = X.shape
        if x_size >= y_size:
            figsize = (6, min(18, 6 * x_size / y_size))
        else:
            figsize = (min(18, 6 * y_size / x_size), 6)

    if cmap is None:
        if X.min(axis=0).min(axis=0) < 0:
            cmap = "RdBu_r"
        else:
            cmap = "hot_r"

    kwargs["cmap"] = cmap
    kwargs = dict(kwargs, interpolation="nearest", aspect="auto")

    if figsize is not False:
        plt.figure(figsize=figsize)

    if ax is None:
        plt.imshow(X, **kwargs)
    else:
        ax.imshow(X, **kwargs)
    plt.grid(None)

    if y is not None:
        y = np.array(y)
        assert all(sorted(y) == y), "This will only work if your row_labels are sorted"

        unik_ys, unik_ys_idx = np.unique(y, return_index=True)
        for u, i in zip(unik_ys, unik_ys_idx):
            plt.hlines(
                i - 0.5,
                0 - 0.5,
                n_cols - 0.5,
                colors="b",
                linestyles="dotted",
                alpha=0.5,
            )
        plt.hlines(
            n_items - 0.5,
            0 - 0.5,
            n_cols - 0.5,
            colors="b",
            linestyles="dotted",
            alpha=0.5,
        )
        plt.yticks(
            unik_ys_idx + np.diff(np.hstack((unik_ys_idx, n_items))) / 2, unik_ys
        )
    elif isinstance(X, pd.DataFrame):
        y_tick_labels = list(X.index)
        plt.yticks(list(range(len(y_tick_labels))), y_tick_labels)

    if col_labels is not None:
        plt.xticks(list(range(len(col_labels))), col_labels)
    else:
        plt.xticks([])

    plt.gca().xaxis.set_tick_params(labeltop=xlabel_top, labelbottom=xlabel_bottom)
    plt.gca().yaxis.set_tick_params(labelleft=ylabel_left, labelright=ylabel_right)

    if return_gcf:
        return plt.gcf()


def plot_life_course(df, grid=False, figsize=3.5, **kwargs):
    if isinstance(
        figsize, (int, float)
    ):  # if figsize is a number, it's a factor of the df size (shape)
        figsize = np.array(df.shape) / figsize
    kwargs["figsize"] = figsize
    heatmap(df.T, **kwargs)
    plt.grid(grid)


def _rgb_array(colors):
    return np.array([to_rgb(c) for c in colors], dtype=float)


def _normalized_values(vals):
    """Map values to [0, 1]: 0 is "off" (background), 1 is "on" (row color)."""
    if np.isnan(vals).all():
        return np.zeros_like(vals)
    lo, hi = min(0.0, np.nanmin(vals)), np.nanmax(vals)
    if not hi > lo:
        return np.zeros_like(vals)
    return np.nan_to_num(np.clip((vals - lo) / (hi - lo), 0, 1), nan=0.0)


def _as_mapping_if_keyed(x):
    """A ``pd.Series`` (colors keyed by name) is used as a mapping, not a sequence."""
    return x.to_dict() if isinstance(x, pd.Series) else x


def _as_color_sequence(x):
    """A single color (``"red"``, ``(1, 0, 0)``) is a sequence of one color."""
    return (x,) if is_color_like(x) else tuple(x)


def _row_foregrounds(varnames, row_colors, default_color):
    row_colors = _as_mapping_if_keyed(row_colors)
    if row_colors is None or (
        not isinstance(row_colors, Mapping) and is_color_like(row_colors)
    ):  # a single color (a name, a hex string, or an RGB(A) tuple)
        return _rgb_array([row_colors or default_color] * len(varnames))
    if isinstance(row_colors, Mapping):
        return _rgb_array([row_colors.get(v, default_color) for v in varnames])
    row_colors = list(row_colors)
    if len(row_colors) != len(varnames):
        raise ValueError(
            f"row_colors has {len(row_colors)} colors, but there are "
            f"{len(varnames)} variables: Give one color per variable, or a "
            "{variable: color} mapping."
        )
    return _rgb_array(row_colors)


def _group_of_var(varnames, row_groups):
    """Return ``(group_of_row, groups)``: the group (or None) of each row, in order."""
    if row_groups is None:
        return [None] * len(varnames), []
    if not isinstance(row_groups, Mapping):
        raise TypeError(
            "row_groups should be a {group_name: iterable_of_variables} mapping, "
            f"not a {type(row_groups).__name__}"
        )
    group_of = {}
    for group, members in row_groups.items():
        for v in [members] if isinstance(members, str) else members:
            group_of.setdefault(v, group)
    return [group_of.get(v) for v in varnames], list(row_groups)


def _row_backgrounds(
    varnames, *, row_backgrounds, row_groups, group_colors, group_tint
):
    n = len(varnames)
    row_backgrounds = _as_color_sequence(row_backgrounds or ("#ffffff",))
    bgs = _rgb_array([c for c, _ in zip(cycle(row_backgrounds), range(n))])
    group_of_row, groups = _group_of_var(varnames, row_groups)
    if groups:
        group_colors = _as_mapping_if_keyed(group_colors)
        if group_colors is None:
            group_colors = DFLT_GROUP_COLORS
        if not isinstance(group_colors, Mapping):
            group_colors = dict(zip(groups, cycle(_as_color_sequence(group_colors))))
        if missing := [g for g in groups if g not in group_colors]:
            raise ValueError(f"group_colors has no color for the groups: {missing}")
        for i, g in enumerate(group_of_row):
            if g is not None:
                tint = np.array(to_rgb(group_colors[g]))
                bgs[i] = (1 - group_tint) * bgs[i] + group_tint * tint
    return bgs, group_of_row


def _major_boundaries(labels, major_every):
    """Indices ``j`` such that a wider gap should be drawn just before column ``j``.

    >>> _major_boundaries([16, 17, 18, 19, 20, 21], 5)
    [4]
    >>> _major_boundaries(['a', 'b', 'c', 'd'], 2)
    [2]
    """
    if not major_every:
        return []
    numeric = all(
        isinstance(x, (int, np.integer))
        or (isinstance(x, (float, np.floating)) and float(x).is_integer())
        for x in labels
    )
    if numeric:
        return [j for j in range(1, len(labels)) if int(labels[j]) % major_every == 0]
    return [j for j in range(1, len(labels)) if j % major_every == 0]


def plot_trajectory(
    df,
    fields=None,
    *,
    title=None,
    ax=None,
    row_colors=None,
    default_color=DFLT_FG_COLOR,
    row_backgrounds=DFLT_ROW_BACKGROUNDS,
    row_groups=None,
    group_colors=None,
    group_tint=DFLT_GROUP_TINT,
    gap_color="white",
    minor_gap=0.6,
    major_gap=2.5,
    major_every=DFLT_MAJOR_EVERY,
    xtick_every=None,
    cell_size=DFLT_CELL_SIZE,
    fontsize=8,
    ylabel_right=False,
):
    """Plot one trajectory (index: ages or years, columns: variables) readably.

    Each variable is a row and each age/year a column. A cell takes its row's
    color in proportion to its value, over the row's background.

    Args:
        df: A trajectory DataFrame: one row per age/year, one column per variable.
        fields: The variables (columns) to show, in order. Defaults to all.
        title: Optional title.
        ax: The axes to draw on. If None, a new figure is made, sized so that
            cells are ``cell_size`` inches and every label fits.
        row_colors: Color of the "on" cells: a single color, a list (one per
            variable), or a ``{variable: color}`` mapping (missing variables get
            ``default_color``). See ``row_colors_of_xls`` to reuse the colors of
            the source spreadsheet.
        default_color: Color of variables that ``row_colors`` doesn't specify.
        row_backgrounds: Background colors, cycled row by row (the alternating
            greys that make it easy to follow a row). Use ``None`` for white.
        row_groups: A ``{group_name: variables}`` mapping. Rows of a group get a
            light tint of its color, and a wider gap separates rows of different
            groups. Keep it to a few (2 to 4) groups.
        group_colors: Colors of the groups: a list (cycled) or a
            ``{group_name: color}`` mapping. Defaults to ``DFLT_GROUP_COLORS``.
        group_tint: How strongly (0 to 1) a group color tints the background.
        gap_color: Color of the gaps between cells.
        minor_gap: Width (points) of the gap between two cells.
        major_gap: Width (points) of the wider gaps: before every age/year
            multiple of ``major_every``, and between row groups.
        major_every: See ``major_gap``. If the index isn't integer-valued, every
            ``major_every``-th column. Falsy: no wider gaps.
        xtick_every: Label every ``xtick_every``-th column. Defaults to every
            column when there's room for it, else to ``major_every``.
        cell_size: Size (inches) of a cell, used when ``ax`` is None.
        fontsize: Font size of the tick labels.
        ylabel_right: Whether to also label the variables on the right.

    Returns:
        The matplotlib axes drawn on.
    """
    if isinstance(fields, slice):
        df = df.loc[:, fields]
    elif fields is not None:
        df = df[[fields] if isinstance(fields, str) else list(fields)]
    varnames, times = list(df.columns), list(df.index)
    n_rows, n_cols = len(varnames), len(times)
    if n_rows == 0 or n_cols == 0:
        raise ValueError(f"Nothing to plot: the trajectory has shape {df.shape}")

    v = _normalized_values(df.T.to_numpy(dtype=float))[..., None]
    fg = _row_foregrounds(varnames, row_colors, default_color)[:, None, :]
    bg, group_of_row = _row_backgrounds(
        varnames,
        row_backgrounds=row_backgrounds,
        row_groups=row_groups,
        group_colors=group_colors,
        group_tint=group_tint,
    )
    img = (1 - v) * bg[:, None, :] + v * fg

    time_labels = [str(t) for t in times]
    if xtick_every is not None and not (
        isinstance(xtick_every, (int, np.integer)) and xtick_every >= 1
    ):
        raise ValueError(f"xtick_every should be a positive int, not {xtick_every!r}")
    auto_xticks = xtick_every is None
    if auto_xticks:
        longest = max(map(len, time_labels))
        fits = longest * fontsize * 0.6 / 72 <= cell_size
        xtick_every = 1 if fits or not major_every else major_every

    if ax is None:
        label_inches = max(map(len, map(str, varnames))) * fontsize * 0.6 / 72
        width = n_cols * cell_size + label_inches * (2 if ylabel_right else 1) + 0.5
        height = n_rows * cell_size + 1.0
        _, ax = plt.subplots(figsize=(width, height), layout="constrained")

    ax.imshow(img, interpolation="nearest", aspect="auto")

    majors = set(_major_boundaries(times, major_every))
    xs = [j - 0.5 for j in range(1, n_cols) if j not in majors]
    ax.vlines(xs, -0.5, n_rows - 0.5, colors=gap_color, linewidth=minor_gap)
    ax.vlines(
        [j - 0.5 for j in sorted(majors)],
        -0.5,
        n_rows - 0.5,
        colors=gap_color,
        linewidth=major_gap,
    )
    group_breaks = {
        i for i in range(1, n_rows) if group_of_row[i] != group_of_row[i - 1]
    }
    ys = [i - 0.5 for i in range(1, n_rows) if i not in group_breaks]
    ax.hlines(ys, -0.5, n_cols - 0.5, colors=gap_color, linewidth=minor_gap)
    ax.hlines(
        [i - 0.5 for i in sorted(group_breaks)],
        -0.5,
        n_cols - 0.5,
        colors=gap_color,
        linewidth=major_gap,
    )

    if auto_xticks and xtick_every > 1 and majors:
        xticks = [j for j in range(n_cols) if j == 0 or j in majors]
    else:
        xticks = list(range(0, n_cols, xtick_every))
    ax.set_xticks(xticks, [time_labels[j] for j in xticks], fontsize=fontsize)
    ax.set_yticks(range(n_rows), [str(v) for v in varnames], fontsize=fontsize)
    ax.tick_params(length=0, labelright=ylabel_right)
    ax.set_xlim(-0.5, n_cols - 0.5)
    ax.set_ylim(n_rows - 0.5, -0.5)
    for spine in ax.spines.values():
        spine.set_visible(False)
    ax.grid(False)
    if title is not None:
        ax.set_title(title)
    return ax


def row_colors_of_xls(xls_filepath, *, sheetname=None, label_col=2, data_col=3):
    """Get a ``{variable: color}`` mapping from the cell colors of a spreadsheet.

    Use it as the ``row_colors`` of ``plot_trajectory`` to color variables as the
    source spreadsheet does. A variable's color is the fill of its label cell, or
    if that has none, the most common fill of its data cells. Variables with no
    solid RGB or indexed (palette) fill are left out, as are theme colors.

    Args:
        xls_filepath: Path to the ``.xlsx`` file.
        sheetname: The sheet to read (default: the active one).
        label_col: The (1-based) column holding the variable names.
        data_col: The (1-based) first column holding the data.
    """
    from collections import Counter
    from openpyxl import load_workbook
    from openpyxl.styles.colors import COLOR_INDEX

    def rgb_of(cell):
        fill = cell.fill
        if fill is None or fill.fill_type != "solid":
            return None
        color = fill.fgColor
        if color.type == "indexed" and color.indexed < len(COLOR_INDEX):
            argb = COLOR_INDEX[color.indexed]
        elif color.type == "rgb" and isinstance(color.rgb, str):
            argb = color.rgb
        else:
            return None  # theme colors (would need the workbook's theme)
        if len(argb) != 8 or argb == "00000000":  # openpyxl's "no color"
            return None
        return "#" + argb[2:].lower()  # Excel ignores the alpha byte

    wb = load_workbook(xls_filepath)
    sheet = wb[sheetname] if sheetname is not None else wb.active
    colors = {}
    for row in sheet.iter_rows(min_col=label_col):
        label = row[0].value
        if not isinstance(label, str) or label in colors:
            continue
        color = rgb_of(row[0])
        if color is None:
            fills = Counter(filter(None, map(rgb_of, row[data_col - label_col :])))
            color = fills.most_common(1)[0][0] if fills else None
        if color is not None:
            colors[label] = color
    return colors


def plot_life(df, fields=None, title=None, ax=None, **plot_kwargs):
    """Plot a trajectory; ``plot_kwargs`` go to ``plot_trajectory``."""
    return plot_trajectory(df, fields, title=title, ax=ax, **plot_kwargs)


def _get_keys(df_store, keys) -> Iterable:
    if keys is None:
        keys = df_store.keys()
    elif callable(keys):
        keys_filt = keys
        keys = filter(keys_filt, df_store.keys())
    assert isinstance(keys, Iterable), "keys should be iterable at this point"
    return keys


def _title_of_key(k, df):
    """The last segment of a store key (``'surveys/B24.xlsx'`` -> ``'B24.xlsx'``)."""
    return str(k).rsplit("/", 1)[-1]


def life_plots(
    df_store,
    fields=None,
    keys=None,
    k_df_to_title=_title_of_key,
    **plot_kwargs,
):
    """Yield the axes of the trajectory plot of each (selected) key of ``df_store``."""
    keys = _get_keys(df_store, keys)
    for k in keys:
        df = df_store[k]
        yield plot_life(df, fields, title=k_df_to_title(k, df), **plot_kwargs)


def write_trajectories_to_file(
    df_store,
    fields=None,
    keys=None,
    fp="test.pdf",
    pil_write_format=None,
    to_pil_image_kwargs=None,
    *,
    k_df_to_title=_title_of_key,
    plot_kwargs=None,
    **pil_save_kwargs,
):
    """Write the trajectory plots of (selected) keys of ``df_store``, one per page.

    ``plot_kwargs`` are passed on to ``plot_trajectory``.
    """
    keys = _get_keys(df_store, keys)
    plot_kwargs = plot_kwargs or {}

    def figs():
        for k in keys:
            df = df_store[k]
            ax = plot_life(df, fields, title=k_df_to_title(k, df), **plot_kwargs)
            fig = ax.figure
            yield fig
            plt.close(fig)

    write_images(figs(), fp, pil_write_format, to_pil_image_kwargs, **pil_save_kwargs)
