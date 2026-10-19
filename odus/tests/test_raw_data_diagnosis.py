"""Tests for the rule-based diagnosis of odus.raw_data_diagnosis."""

import os

import numpy as np
import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill

from odus.raw_data_diagnosis import (
    DFLT_RULES,
    Rule,
    diagnose_dir,
    diagnose_file,
    diagnosis_table,
    print_diagnosis,
    select_rules,
)

FILL = PatternFill("solid", fgColor="FFFFFF00")


def _write_sheet(fp, rows, colored=()):
    wb = Workbook()
    sh = wb.active
    for row in rows:
        sh.append(row)
    for coordinate in colored:
        sh[coordinate].fill = FILL
    wb.save(fp)
    return fp


@pytest.fixture
def clean_rows():
    return [
        [None, "INTERVIEW ID NUMBER", "X01", None, None],
        [0, "YEAR", 2000, 2001, 2002],
        [1, "AGE", 20, 21, 22],
        [3, "HEROIN", 0, 1, 1],
        [4, "WORK", 1, 1, 0],
    ]


def _problems(fp, **kwargs):
    return {(p.rule, p.coordinate) for p in diagnose_file(fp, **kwargs)}


def test_clean_sheet_has_no_problems(tmp_path, clean_rows):
    assert _problems(_write_sheet(tmp_path / "x.xlsx", clean_rows)) == set()


def test_value_and_header_rules(tmp_path, clean_rows):
    rows = clean_rows
    rows[1][4] = 2003  # year skips 2002
    rows[2][3] = 25  # age doesn't follow 20 (and doesn't match its year)
    rows[3][3] = 16  # not binary
    rows[4][4] = "x"  # not binary
    rows.append([5, "HEROIN", 0, 0, 0])  # duplicate variable
    problems = _problems(_write_sheet(tmp_path / "x.xlsx", rows))
    assert problems == {
        ("year_not_consecutive", "E2"),
        ("age_not_consecutive", "D3"),
        ("age_not_consecutive", "E3"),
        ("age_year_mismatch", "D3"),
        ("age_year_mismatch", "E3"),  # 2003 - 22 (the year skipped 2002)
        ("not_binary", "D4"),
        ("not_binary", "E5"),
        ("duplicate_variable", "B6"),
    }


def test_color_rules(tmp_path, clean_rows):
    rows = clean_rows
    rows[3][2] = None  # C4 empty...
    colored = ["C4", "D4", "E4", "C5", "D5"]  # ...but colored; E5 filled, uncolored
    problems = _problems(_write_sheet(tmp_path / "x.xlsx", rows, colored))
    assert problems == {
        ("empty_but_colored", "C4"),
        ("filled_but_uncolored", "E5"),
    }


def test_filled_but_uncolored_is_silent_when_a_sheet_has_no_colors(
    tmp_path, clean_rows
):
    fp = _write_sheet(tmp_path / "x.xlsx", clean_rows)
    rules = select_rules(include=["filled_but_uncolored"])
    assert _problems(fp, rules=rules) == set()


def test_messages_explain_the_problem(tmp_path, clean_rows):
    clean_rows[3][3] = 16
    (p,) = diagnose_file(_write_sheet(tmp_path / "x.xlsx", clean_rows))
    assert p.filename == "x.xlsx"
    assert p.message == "D4 (HEROIN) should be 0 or 1, but is 16"


def test_turning_rules_on_and_off_and_adding_one(tmp_path, clean_rows):
    clean_rows[3][3] = 16
    fp = _write_sheet(tmp_path / "x.xlsx", clean_rows)
    assert _problems(fp, rules=select_rules(exclude=["not_binary"])) == set()

    def flag_work(view):
        mask = view.false_mask()
        mask[:, view.label_col] = (view.labels == "WORK").to_numpy()
        return mask

    no_work = Rule("no_work", flag_work, "{coordinate}: {value} isn't tracked anymore")
    rules = select_rules(include=["not_binary"], extra=[no_work])
    assert _problems(fp, rules=rules) == {("not_binary", "D4"), ("no_work", "B5")}
    assert set(DFLT_RULES) >= {"not_binary", "empty_but_colored"}  # untouched

    def over_21(view):
        return view.data_mask & (view.numeric()[view.age_row] > 21)

    rules = select_rules(include=[], extra=[Rule("old", over_21, "{coordinate}")])
    assert _problems(fp, rules=rules) == {("old", "E4"), ("old", "E5")}

    with pytest.raises(KeyError, match="Unknown rule"):
        select_rules(exclude=["no_such_rule"])


def test_a_rule_returning_a_wrong_shape_is_reported(tmp_path, clean_rows):
    fp = _write_sheet(tmp_path / "x.xlsx", clean_rows)
    bad = Rule("bad", lambda view: np.zeros((1, 1), dtype=bool))
    with pytest.raises(ValueError, match="'bad' returned a mask of shape"):
        list(diagnose_file(fp, rules=[bad]))


def test_dir_table_and_print(tmp_path, clean_rows, capsys):
    _write_sheet(tmp_path / "clean.xlsx", clean_rows)
    clean_rows[3][3] = 16
    _write_sheet(tmp_path / "bad.xlsx", clean_rows)
    table = diagnosis_table(diagnose_dir(tmp_path))
    assert list(table.columns) == ["filename", "rule", "coordinate", "value", "message"]
    assert table.filename.tolist() == ["bad.xlsx"]

    print_diagnosis(str(tmp_path))
    out = capsys.readouterr().out
    assert "bad.xlsx" in out and "clean.xlsx" not in out
    assert "not_binary:" in out and "D4 (HEROIN) should be 0 or 1" in out

    os.remove(tmp_path / "bad.xlsx")
    print_diagnosis(str(tmp_path))
    assert "No problems found." in capsys.readouterr().out


def test_labels_in_column_a_with_a_blank_first_variable_row(tmp_path):
    rows = [
        ["INTERVIEW ID NUMBER", "X01", None, None],
        ["YEAR", 2000, 2001, 2002],
        ["AGE", 20, 21, 22],
        [None, None, None, None],
        ["HEROIN", 0, 7, 1],
    ]
    problems = _problems(_write_sheet(tmp_path / "x.xlsx", rows))
    assert problems == {("not_binary", "C5")}


def test_a_bad_first_column_is_the_one_flagged(tmp_path, clean_rows):
    clean_rows[2][2] = 19  # ages 19, 21, 22 for years 2000, 2001, 2002
    problems = _problems(_write_sheet(tmp_path / "x.xlsx", clean_rows))
    assert ("age_year_mismatch", "C3") in problems
    assert ("age_year_mismatch", "D3") not in problems


def test_unreadable_sheet_is_reported_and_the_dir_run_goes_on(tmp_path, clean_rows):
    _write_sheet(tmp_path / "a_short.xlsx", clean_rows[:2])
    clean_rows[3][3] = 16
    _write_sheet(tmp_path / "b_bad.xlsx", clean_rows)
    table = diagnosis_table(diagnose_dir(tmp_path))
    assert table[["filename", "rule"]].values.tolist() == [
        ["a_short.xlsx", "unreadable_sheet"],
        ["b_bad.xlsx", "not_binary"],
    ]


def test_header_colors_or_white_fills_dont_count_as_colored_data(tmp_path, clean_rows):
    white = PatternFill("solid", fgColor="FFFFFFFF")
    fp = tmp_path / "x.xlsx"
    _write_sheet(fp, clean_rows, colored=["A1"])
    assert _problems(fp) == set()  # only a header cell is colored
    clean_rows[3][2] = None
    wb = Workbook()
    sh = wb.active
    for row in clean_rows:
        sh.append(row)
    sh["C4"].fill = white
    wb.save(fp)
    assert _problems(fp) == set()  # an empty, white-filled cell


def test_duplicates_ignore_case_and_spaces(tmp_path, clean_rows):
    clean_rows.append([5, "heroin ", 0, 0, 0])
    problems = _problems(_write_sheet(tmp_path / "x.xlsx", clean_rows))
    assert problems == {("duplicate_variable", "B6")}


def test_rules_conveniences(tmp_path, clean_rows):
    clean_rows[3][3] = 16
    fp = tmp_path / "x.xlsx"
    _write_sheet(fp, clean_rows)
    assert list(select_rules(include="not_binary")) == ["not_binary"]
    bare = Rule("bare", lambda view: view.vals.map(lambda x: x == 16).to_numpy())
    (p,) = diagnose_file(fp, rules=bare)  # a single rule, no message
    assert p.message == "D4 (HEROIN): 16"
