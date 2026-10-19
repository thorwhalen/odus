"""Tests for odus.plot_utils."""

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import pytest
from matplotlib.colors import to_rgb

from odus.plot_utils import (
    plot_life,
    plot_trajectory,
    row_colors_of_xls,
    stacked_trajectories,
    write_trajectories_to_file,
)


@pytest.fixture
def traj():
    ages = list(range(18, 30))
    return pd.DataFrame(
        {
            "HEROIN": [0] * 6 + [1] * 6,
            "WORK": [1] * 12,
            "SPOUSE": [0] * 12,
        },
        index=pd.Index(ages, name="age"),
    )


@pytest.fixture(autouse=True)
def _close_figs():
    yield
    plt.close("all")


def _img(ax):
    (im,) = ax.get_images()
    return np.asarray(im.get_array())


def test_cells_take_row_color_over_alternating_background(traj):
    ax = plot_trajectory(
        traj,
        ["HEROIN", "SPOUSE", "WORK"],
        row_colors={"HEROIN": "red"},
        default_color="black",
        row_backgrounds=("white", "#cccccc"),
    )
    img = _img(ax)
    assert img.shape == (3, 12, 3)  # one row per variable, one column per age
    assert np.allclose(img[0, 0], to_rgb("white"))  # HEROIN off: row 0 background
    assert np.allclose(img[0, -1], to_rgb("red"))  # HEROIN on: its own color
    assert np.allclose(img[1, 0], to_rgb("#cccccc"))  # SPOUSE off: alternated grey
    assert np.allclose(img[2, 0], to_rgb("black"))  # WORK on: default color
    assert [t.get_text() for t in ax.get_yticklabels()] == ["HEROIN", "SPOUSE", "WORK"]


def test_major_gaps_before_multiples_of_major_every(traj):
    ax = plot_trajectory(traj, minor_gap=0.5, major_gap=3, major_every=5)
    widths = {}
    for coll in ax.collections:
        for seg in coll.get_segments():
            (x0, _), (x1, _) = seg
            if x0 == x1:  # vertical: a gap between two columns
                widths[x0] = coll.get_linewidth()[0]
    # ages 18..29: wider gaps just before ages 20 and 25 (columns 2 and 7)
    assert {x for x, w in widths.items() if w == 3} == {1.5, 6.5}
    assert len(widths) == len(traj) - 1


def test_row_groups_tint_background_and_separate_groups(traj):
    ax = plot_trajectory(
        traj,
        row_groups={"drugs": ["HEROIN"]},
        group_colors={"drugs": "blue"},
        group_tint=0.5,
        row_backgrounds=("white",),
    )
    img = _img(ax)
    assert np.allclose(img[0, 0], [0.5, 0.5, 1.0])  # tinted towards blue
    assert np.allclose(img[2, 0], [1, 1, 1])  # not in a group: untouched
    horizontal_widths = {
        seg[0][1]: coll.get_linewidth()[0]
        for coll in ax.collections
        for seg in coll.get_segments()
        if seg[0][1] == seg[1][1]
    }
    assert horizontal_widths[0.5] > horizontal_widths[1.5]  # group boundary


def test_fields_select_and_order_variables(traj):
    ax = plot_life(traj, fields=["SPOUSE", "HEROIN"], title="t")
    assert [t.get_text() for t in ax.get_yticklabels()] == ["SPOUSE", "HEROIN"]
    assert ax.get_title() == "t"


def test_labels_fit_in_the_figure(traj):
    traj = traj.rename(columns={"WORK": "A VERY LONG VARIABLE NAME, TO BE SURE"})
    ax = plot_trajectory(traj)
    fig = ax.figure
    fig.canvas.draw()
    fig_box = fig.bbox
    for label in ax.get_yticklabels() + ax.get_xticklabels():
        box = label.get_window_extent()
        assert box.x0 >= fig_box.x0 - 1 and box.x1 <= fig_box.x1 + 1


def test_an_rgb_tuple_is_one_color_not_one_per_variable(traj):
    img = _img(plot_trajectory(traj, ["WORK"] * 3, row_colors=(1.0, 0.0, 0.0)))
    assert np.allclose(img[:, 0], [1.0, 0.0, 0.0])


def test_row_colors_length_mismatch_is_an_informative_error(traj):
    with pytest.raises(ValueError, match="one color per variable"):
        plot_trajectory(traj, row_colors=["red"])


def test_row_colors_of_xls(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    wb = Workbook()
    sh = wb.active
    sh.append([None, "label", 2000, 2001])
    sh.append([None, "HEROIN", 1, 0])
    sh.append([None, "WORK", 1, 1])
    sh.append([None, "SPOUSE", 0, 0])
    sh["B2"].fill = PatternFill("solid", fgColor="FFFF0000")  # label cell
    sh["C3"].fill = PatternFill("solid", fgColor="FF00FF00")  # data cells
    sh["D3"].fill = PatternFill("solid", fgColor="FF00FF00")
    fp = tmp_path / "t.xlsx"
    wb.save(fp)
    assert row_colors_of_xls(fp) == {"HEROIN": "#ff0000", "WORK": "#00ff00"}


def test_write_trajectories_to_file(traj, tmp_path):
    store = {"surveys/A.xlsx": traj, "surveys/B.xlsx": traj}
    fp = tmp_path / "out.pdf"
    n_figs = len(plt.get_fignums())
    write_trajectories_to_file(store, fp=str(fp), plot_kwargs={"row_backgrounds": None})
    assert fp.stat().st_size > 0
    assert len(plt.get_fignums()) == n_figs  # figures are closed after writing


def test_explicit_xtick_every_is_honored(traj):
    ax = plot_trajectory(traj, xtick_every=2)  # major gaps at 20 and 25
    assert [t.get_text() for t in ax.get_xticklabels()] == [
        "18", "20", "22", "24", "26", "28"
    ]  # fmt: skip
    with pytest.raises(ValueError, match="positive int"):
        plot_trajectory(traj, xtick_every=0)


def test_series_row_colors_are_keyed_by_variable(traj):
    colors = pd.Series({"SPOUSE": "green", "WORK": "red", "HEROIN": "blue"})
    img = _img(plot_trajectory(traj, ["HEROIN", "WORK"], row_colors=colors))
    assert np.allclose(img[0, -1], to_rgb("blue"))
    assert np.allclose(img[1, -1], to_rgb("red"))


def test_fields_as_a_slice_or_a_single_name(traj):
    names = lambda ax: [t.get_text() for t in ax.get_yticklabels()]
    assert names(plot_life(traj, slice(None))) == list(traj.columns)
    assert names(plot_life(traj, "WORK")) == ["WORK"]


def test_single_colors_and_missing_group_colors(traj):
    img = _img(
        plot_trajectory(
            traj,
            ["SPOUSE"],
            row_backgrounds="red",
            row_groups={"g": ["SPOUSE"]},
            group_colors="blue",
            group_tint=0.5,
        )
    )
    assert np.allclose(img[0, 0], [0.5, 0.0, 0.5])
    with pytest.raises(ValueError, match="no color for the groups"):
        plot_trajectory(
            traj,
            row_groups={"g1": ["WORK"], "g2": ["SPOUSE"]},
            group_colors={"g1": "b"},
        )


def test_row_colors_of_xls_reads_six_hex_and_indexed_fills(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    from openpyxl.styles.colors import Color

    wb = Workbook()
    sh = wb.active
    sh.append([None, "SIXHEX", 1])
    sh.append([None, "INDEXED", 1])
    sh.append([None, "THEME", 1])
    sh["B1"].fill = PatternFill("solid", fgColor="FF0000")  # stored as 00FF0000
    sh["B2"].fill = PatternFill("solid", fgColor=Color(indexed=13))  # yellow
    sh["B3"].fill = PatternFill("solid", fgColor=Color(theme=4))
    fp = tmp_path / "t.xlsx"
    wb.save(fp)
    assert row_colors_of_xls(fp) == {"SIXHEX": "#ff0000", "INDEXED": "#ffff00"}


def test_store_keys_without_a_folder_make_titles(traj):
    from odus.plot_utils import life_plots

    (ax,) = life_plots({"B24.xlsx": traj})
    assert ax.get_title() == "B24.xlsx"


@pytest.fixture
def two_trajectories(traj):
    late = traj.copy()
    late.index = late.index + 10  # ages 28..39
    late["HEROIN"] = [0] * 3 + [1] * 9  # first use at 31 (first use at 24 in traj)
    traj = traj.copy()
    traj.attrs["year_of_age"] = {a: a + 1980 for a in traj.index}
    late.attrs["year_of_age"] = {a: a + 1960 for a in late.index}
    return {"s/A.xlsx": traj, "s/B.xlsx": late}


def _extents(fig):
    return [tuple(map(float, ax.get_images()[0].get_extent()[:2])) for ax in fig.axes]


@pytest.mark.parametrize(
    "x_align, expected_extents, xlim",
    [
        ("age", [(17.5, 29.5), (27.5, 39.5)], (17.5, 39.5)),
        ("year", [(1997.5, 2009.5), (1987.5, 1999.5)], (1987.5, 2009.5)),
        ("HEROIN", [(-6.5, 5.5), (-3.5, 8.5)], (-6.5, 8.5)),
    ],
)
def test_stacked_trajectories_align(two_trajectories, x_align, expected_extents, xlim):
    fig = stacked_trajectories(two_trajectories, x_align=x_align)
    assert [ax.get_title() for ax in fig.axes] == ["A.xlsx", "B.xlsx"]
    assert _extents(fig) == expected_extents
    assert all(ax.get_xlim() == xlim for ax in fig.axes)  # one common x axis
    assert fig.axes[-1].get_xlabel() == (
        x_align if x_align in ("age", "year") else "years since first HEROIN"
    )


def test_stacked_trajectories_leave_out_what_cant_be_aligned(two_trajectories):
    with pytest.warns(UserWarning, match=r"Left out 1 .*'SPOUSE'.*\['s/A.xlsx'\]"):
        two_trajectories["s/B.xlsx"]["SPOUSE"] = 1
        fig = stacked_trajectories(
            two_trajectories, ["SPOUSE", "WORK"], x_align="SPOUSE"
        )
    assert [ax.get_title() for ax in fig.axes] == ["B.xlsx"]
    with (
        pytest.raises(ValueError, match="No trajectory to plot"),
        pytest.warns(UserWarning),
    ):
        stacked_trajectories(two_trajectories, x_align="NOT A VARIABLE")


def test_stacked_trajectories_with_a_function_and_bad_alignments(two_trajectories):
    fig = stacked_trajectories(
        two_trajectories, keys=["s/A.xlsx"], x_align=lambda df: df.index - 18
    )
    assert _extents(fig) == [(-0.5, 11.5)]
    with pytest.raises(ValueError, match="go up by 1"):
        stacked_trajectories(two_trajectories, x_align=lambda df: df.index * 2)
    del two_trajectories["s/A.xlsx"].attrs["year_of_age"]
    with pytest.raises(ValueError, match="needs the years"):
        stacked_trajectories(two_trajectories, x_align="year")


def test_df_store_trajectories_stack_by_year():
    import os
    import warnings

    from odus import data_dir
    from odus.dacc import DfStore

    store = DfStore(os.path.join(data_dir, "surveys") + os.sep)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        fig = stacked_trajectories(
            store, ["HEROIN"], keys=["A01.xlsx", "A60.xlsx"], x_align="year"
        )
    a01 = store["A01.xlsx"]
    first_year = a01.attrs["year_of_age"][a01.index[0]]
    assert fig.axes[0].get_images()[0].get_extent()[0] == first_year - 0.5
